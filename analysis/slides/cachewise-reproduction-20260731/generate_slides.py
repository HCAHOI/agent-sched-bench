#!/usr/bin/env python3
"""Generate the editable CacheWise transfer/failure group-meeting deck.

Run from the repository root with python-pptx available, for example:
  uv run --with python-pptx python analysis/slides/cachewise-reproduction-20260731/generate_slides.py
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
OUT = HERE / "cachewise-reproduction-20260731.pptx"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
MONO = "Liberation Mono"

BG = "0B1020"
PANEL = "131B2E"
PANEL2 = "19243A"
WHITE = "F4F7FB"
MUTED = "9FB0C6"
FAINT = "64748B"
CYAN = "2DD4BF"
BLUE = "60A5FA"
GREEN = "34D399"
YELLOW = "FBBF24"
CORAL = "FB7185"
RED = "F87171"
PURPLE = "A78BFA"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def load_result(name: str) -> dict:
    path = ROOT / "analysis" / "results" / name / "result.json"
    return json.loads(path.read_text(encoding="utf-8"))


SWE = load_result("cachewise-swe-reproduction-20260731")
PROGRESS = load_result("cachewise-progress-counter-20260731")
SWE_SYNTH = load_result("cachewise-swe-synthetic-c32-20260731")
TB_SYNTH = load_result("cachewise-tb-synthetic-c32-20260731")
TIMEOUT = load_result("cachewise-tb-timeout-stability-20260731")


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = WHITE,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = FONT,
    margin: float = 0.03,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    return shape


def add_rich(
    slide,
    runs: Iterable[tuple[str, str, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = PANEL,
    line: str | None = None,
    radius: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def line(slide, x1, y1, x2, y2, *, color=FAINT, width=2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def circle(slide, x, y, d, *, fill=CYAN, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color or fill)
    return shape


def badge(slide, text, x, y, w, *, fill=PANEL2, color=WHITE):
    rect(slide, x, y, w, 0.34, fill=fill, line=fill)
    add_text(
        slide,
        text,
        x + 0.08,
        y + 0.015,
        w - 0.16,
        0.26,
        size=10.5,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def title(slide, text: str, page: int, *, kicker: str | None = None):
    if kicker:
        add_text(
            slide, kicker.upper(), 0.6, 0.22, 6.0, 0.25, size=9.5, color=CYAN, bold=True
        )
    add_text(slide, text, 0.58, 0.48, 12.1, 0.72, size=27, bold=True)
    line(slide, 0.6, 1.14, 12.73, 1.14, color=PANEL2, width=1.2)
    add_text(
        slide,
        f"{page:02d}",
        12.45,
        0.22,
        0.3,
        0.25,
        size=9.5,
        color=FAINT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def footer(slide, text: str):
    line(slide, 0.6, 7.08, 12.73, 7.08, color=PANEL2, width=0.8)
    add_text(slide, text, 0.6, 7.12, 12.0, 0.22, size=8.2, color=FAINT)


def stat_card(slide, x, y, w, value, label, *, color=CYAN, sub=None, value_size=23):
    rect(slide, x, y, w, 1.05, fill=PANEL, line=PANEL2)
    add_text(
        slide,
        value,
        x + 0.16,
        y + 0.12,
        w - 0.32,
        0.38,
        size=value_size,
        color=color,
        bold=True,
    )
    add_text(
        slide,
        label,
        x + 0.16,
        y + 0.52,
        w - 0.32,
        0.26,
        size=12.5,
        color=WHITE,
        bold=True,
    )
    if sub:
        add_text(slide, sub, x + 0.16, y + 0.78, w - 0.32, 0.18, size=9.5, color=MUTED)


def bullet_list(slide, items, x, y, w, h, *, size=16, color=WHITE, bullet_color=CYAN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(8)
    return box


def hbars(slide, labels, values, x, y, w, h, *, colors=None, unit="s", lower=True):
    max_value = max(values) * 1.08
    row_h = h / len(values)
    for i, (label, value) in enumerate(zip(labels, values, strict=True)):
        yy = y + i * row_h
        add_text(
            slide,
            label,
            x,
            yy + 0.03,
            1.0,
            row_h - 0.05,
            size=12.5,
            color=MUTED,
            bold=True,
        )
        bar_x = x + 1.05
        bar_w = (w - 2.2) * value / max_value
        color = colors[i] if colors else BLUE
        rect(
            slide,
            bar_x,
            yy + 0.1,
            max(bar_w, 0.03),
            row_h - 0.22,
            fill=color,
            line=color,
            radius=False,
        )
        add_text(
            slide,
            f"{value:.3f} {unit}",
            x + w - 1.1,
            yy + 0.03,
            1.05,
            row_h - 0.05,
            size=12,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
    add_text(
        slide,
        "↓ lower is better" if lower else "↑ higher is better",
        x + 1.05,
        y + h + 0.02,
        1.8,
        0.2,
        size=9.5,
        color=MUTED,
    )


def ci_axis(slide, x, y, w, lo, point, hi, *, min_v, max_v, label, color=CORAL):
    def pos(value):
        return x + (value - min_v) / (max_v - min_v) * w

    line(slide, x, y, x + w, y, color=FAINT, width=1.5)
    zero_x = pos(0)
    line(slide, zero_x, y - 0.26, zero_x, y + 0.26, color=WHITE, width=1.4)
    line(slide, pos(lo), y, pos(hi), y, color=color, width=5)
    line(slide, pos(lo), y - 0.11, pos(lo), y + 0.11, color=color, width=2)
    line(slide, pos(hi), y - 0.11, pos(hi), y + 0.11, color=color, width=2)
    circle(slide, pos(point) - 0.075, y - 0.075, 0.15, fill=color)
    add_text(
        slide,
        label,
        x,
        y + 0.18,
        w,
        0.3,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "better",
        x,
        y - 0.42,
        zero_x - x - 0.05,
        0.2,
        size=9,
        color=GREEN,
        align=PP_ALIGN.RIGHT,
    )
    add_text(
        slide, "worse", zero_x + 0.05, y - 0.42, x + w - zero_x, 0.2, size=9, color=RED
    )


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    return slide


def generate() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "CacheWise reproduction and failure analysis"
    prs.core_properties.subject = "Group meeting, 15–20 minutes"
    prs.core_properties.author = "agent-sched-bench"

    # 1 — title
    s = new_slide(prs)
    add_text(
        s,
        "CACHEWISE × AGENT TRACES",
        0.65,
        0.45,
        4.5,
        0.28,
        size=11,
        color=CYAN,
        bold=True,
    )
    add_text(
        s,
        "CacheWise 的核心 predictor\n没有稳定迁移到我们的 agent traces",
        0.62,
        0.92,
        8.2,
        1.55,
        size=31,
        bold=True,
    )
    add_text(
        s,
        "论文复现 · causal progress · synthetic concurrency · task stability",
        0.66,
        2.68,
        7.3,
        0.38,
        size=17,
        color=MUTED,
    )
    rect(s, 9.28, 0.72, 3.38, 2.63, fill=PANEL, line=PANEL2)
    add_text(s, "一句话结论", 9.55, 0.98, 2.7, 0.35, size=13, color=CYAN, bold=True)
    add_text(
        s,
        "metadata 有信号，\n但收益被少数极端任务支配；\n离可部署 policy 还差 task robustness 与真实 KV pressure。",
        9.52,
        1.45,
        2.75,
        1.45,
        size=17,
        bold=True,
    )
    stat_card(
        s,
        0.66,
        4.25,
        3.75,
        "+0.587 s",
        "SWE · C100 更差",
        color=RED,
        sub="95% CI [+0.120, +1.188]",
    )
    stat_card(
        s,
        4.76,
        4.25,
        3.75,
        "−72.299 s",
        "TB · aggregate 更好",
        color=GREEN,
        sub="但 single-task fragile",
    )
    stat_card(
        s,
        8.86,
        4.25,
        3.75,
        "+4.012 s",
        "Timeout · delete-one 翻转",
        color=RED,
        sub="mixed-integer-programming",
    )
    add_text(
        s,
        "16 slides · 15–20 min · development-only evidence",
        0.68,
        6.55,
        5.2,
        0.28,
        size=11,
        color=FAINT,
    )
    footer(
        s,
        "Paper: CacheWise, arXiv:2606.16824v1 · Local evidence: analysis/results/cachewise-*20260731/",
    )

    # 2 — workload problem
    s = new_slide(prs)
    title(
        s,
        "Coding agent 把一次用户任务变成“LLM ↔ tool”的长闭环",
        2,
        kicker="Paper problem",
    )
    nodes = [
        (0.75, "用户任务", BLUE),
        (2.7, "LLM", PURPLE),
        (4.55, "tool", YELLOW),
        (6.4, "LLM", PURPLE),
        (8.25, "tool", YELLOW),
    ]
    for i, (x, label, color) in enumerate(nodes):
        rect(s, x, 1.75, 1.35, 0.72, fill=PANEL, line=color)
        add_text(
            s,
            label,
            x + 0.05,
            1.92,
            1.25,
            0.3,
            size=16,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if i < len(nodes) - 1:
            add_text(
                s,
                "→",
                x + 1.43,
                1.88,
                0.45,
                0.32,
                size=20,
                color=MUTED,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    add_text(
        s,
        "tool 运行时，LLM 不占算力；但会话 KV prefix 仍驻留 GPU，等待下一轮复用",
        1.55,
        2.78,
        8.6,
        0.44,
        size=18,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    rect(s, 10.25, 1.45, 2.25, 2.2, fill=PANEL, line=CORAL)
    add_text(
        s,
        "GPU HBM",
        10.5,
        1.68,
        1.75,
        0.3,
        size=14,
        color=CORAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    for i, c in enumerate([BLUE, PURPLE, CYAN, YELLOW]):
        rect(s, 10.62, 2.18 + i * 0.28, 1.48, 0.2, fill=c, line=c, radius=False)
    add_text(
        s,
        "多个长 session\n竞争有限 KVCache",
        10.42,
        3.18,
        1.9,
        0.52,
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    stat_card(
        s,
        0.75,
        4.55,
        3.6,
        "36 min",
        "CATraces median session",
        color=CYAN,
        sub="tail > 2.6 h",
    )
    stat_card(
        s,
        4.86,
        4.55,
        3.6,
        "20×",
        "tool-triggered turns",
        color=YELLOW,
        sub="median vs user-triggered",
    )
    stat_card(
        s,
        8.97,
        4.55,
        3.6,
        "21×",
        "prefill / decode ratio",
        color=PURPLE,
        sub="vs chatbot workloads",
    )
    footer(
        s,
        "CacheWise paper §3, Figures 2–8. Paper statistics are CATraces claims, not measurements from our traces.",
    )

    # 3 — ranking insight
    s = new_slide(prs)
    title(
        s, "真正的 eviction 问题不是“多久”，而是“谁最晚回来”", 3, kicker="Core insight"
    )
    add_text(
        s,
        "memory pressure at time t",
        0.75,
        1.45,
        2.9,
        0.32,
        size=14,
        color=CORAL,
        bold=True,
    )
    line(s, 1.15, 2.0, 11.8, 2.0, color=FAINT, width=1.2)
    line(s, 2.35, 1.7, 2.35, 5.2, color=WHITE, width=1.6)
    add_text(
        s,
        "now",
        2.08,
        1.4,
        0.6,
        0.25,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    sessions = [
        ("S₁", 2, BLUE, "soon"),
        ("S₂", 40, CORAL, "furthest"),
        ("S₃", 9, YELLOW, "later"),
    ]
    for i, (name, remaining, color, tag) in enumerate(sessions):
        yy = 2.45 + i * 0.85
        add_text(s, name, 0.85, yy - 0.08, 0.55, 0.3, size=16, color=color, bold=True)
        rect(s, 1.55, yy, 0.8, 0.18, fill=PANEL2, line=PANEL2, radius=False)
        width = 0.32 + 7.2 * remaining / 40
        rect(s, 2.35, yy, width, 0.18, fill=color, line=color, radius=False)
        add_text(
            s,
            f"remaining = {remaining}s",
            2.55 + width,
            yy - 0.1,
            1.65,
            0.3,
            size=12,
            color=color,
            bold=True,
        )
        if tag == "furthest":
            badge(s, "evict this KV", 9.98, yy - 0.08, 1.45, fill=CORAL, color=BG)
    rect(s, 0.82, 5.15, 5.6, 1.02, fill=PANEL, line=PANEL2)
    add_text(
        s, "Belady-like target", 1.08, 5.36, 1.7, 0.28, size=12, color=CYAN, bold=True
    )
    add_text(
        s,
        "j* = arg max  τⱼ(t)",
        2.8,
        5.28,
        3.1,
        0.42,
        size=23,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    rect(s, 6.8, 5.15, 5.72, 1.02, fill=PANEL, line=PANEL2)
    add_text(
        s,
        "CacheWise 的关键简化",
        7.08,
        5.32,
        2.25,
        0.3,
        size=13,
        color=YELLOW,
        bold=True,
    )
    add_text(
        s,
        "不必精确预测 duration；只要把“最晚回来”的 session 排到第一。",
        9.2,
        5.25,
        2.92,
        0.55,
        size=16,
        color=WHITE,
        bold=True,
    )
    footer(
        s,
        "CacheWise paper §4–5.2. τ means time to next KV reuse; predictive eviction targets relative ordering.",
    )

    # 4 — two mechanisms
    s = new_slide(prs)
    title(s, "CacheWise 用两个互补机制减少 KV thrashing", 4, kicker="System design")
    rect(s, 0.72, 1.45, 5.85, 4.85, fill=PANEL, line=BLUE)
    badge(s, "① PREFIX-AWARE SCHEDULING", 1.02, 1.72, 2.75, fill=BLUE, color=BG)
    add_text(
        s,
        "先服务需要新增 KV block 最少的 request",
        1.02,
        2.28,
        5.1,
        0.52,
        size=20,
        bold=True,
    )
    add_text(
        s,
        "aᵢ(t) = dᵢ − kᵢ(t)",
        1.05,
        3.02,
        2.8,
        0.45,
        size=24,
        color=BLUE,
        bold=True,
        font=MONO,
    )
    for i, (label, blocks, color) in enumerate(
        [("A", 2, GREEN), ("B", 7, CORAL), ("C", 4, YELLOW)]
    ):
        yy = 3.72 + i * 0.57
        add_text(s, label, 1.08, yy, 0.4, 0.28, size=13, color=color, bold=True)
        for j in range(blocks):
            rect(
                s,
                1.62 + j * 0.34,
                yy + 0.03,
                0.25,
                0.22,
                fill=color,
                line=color,
                radius=False,
            )
        add_text(s, f"+{blocks} blocks", 4.15, yy, 1.35, 0.28, size=12, color=MUTED)
    add_text(
        s,
        "结果：少动别人的 prefix，也近似 shortest-job-first",
        1.02,
        5.64,
        5.0,
        0.38,
        size=14,
        color=MUTED,
    )
    rect(s, 6.78, 1.45, 5.83, 4.85, fill=PANEL, line=CORAL)
    badge(s, "② PREDICTIVE EVICTION", 7.08, 1.72, 2.55, fill=CORAL, color=BG)
    add_text(
        s,
        "真的要 evict 时，优先释放预计最晚复用的 KV",
        7.08,
        2.28,
        5.05,
        0.55,
        size=20,
        bold=True,
    )
    add_text(
        s,
        "tool_name + tool_args + elapsed",
        7.1,
        3.06,
        4.8,
        0.38,
        size=18,
        color=CORAL,
        bold=True,
        font=MONO,
    )
    add_text(
        s,
        "↓",
        9.2,
        3.48,
        0.7,
        0.32,
        size=20,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    rect(s, 7.66, 3.9, 3.8, 0.72, fill=PANEL2, line=CORAL)
    add_text(
        s,
        "E[ remaining | survived to elapsed ]",
        7.82,
        4.1,
        3.48,
        0.28,
        size=15,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=MONO,
    )
    add_text(
        s,
        "vLLM prototype ≈2,500 Python LoC · 每 3 engine iterations 重建 eviction heap",
        7.08,
        5.24,
        5.05,
        0.55,
        size=13.5,
        color=MUTED,
    )
    footer(
        s,
        "CacheWise paper §5.1–5.3. The two mechanisms are independent in the paper ablation.",
    )

    # 5 — predictor pipeline
    s = new_slide(prs)
    title(
        s,
        "Predictor 是 survival estimator，不是普通 duration regression",
        5,
        kicker="Necessary details",
    )
    steps = [
        ("outer args", '{"command":"pytest…"}', BLUE),
        ("TF-IDF", "≤5,000 terms", PURPLE),
        ("KMeans", "C=20 / 50 / 100", YELLOW),
        ("history", "durations in cluster", CYAN),
        ("survive", "keep D > elapsed", GREEN),
        ("rank", "max E[D−e | D>e]", CORAL),
    ]
    for i, (head, body, color) in enumerate(steps):
        x = 0.65 + i * 2.05
        rect(s, x, 1.7, 1.67, 1.22, fill=PANEL, line=color)
        add_text(
            s,
            head,
            x + 0.12,
            1.9,
            1.43,
            0.28,
            size=14,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            s,
            body,
            x + 0.1,
            2.28,
            1.47,
            0.42,
            size=11.5,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            font=MONO,
        )
        if i < len(steps) - 1:
            add_text(
                s,
                "→",
                x + 1.7,
                2.15,
                0.34,
                0.3,
                size=18,
                color=MUTED,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    rect(s, 1.12, 3.52, 5.45, 2.25, fill=PANEL, line=PANEL2)
    add_text(s, "小例子", 1.4, 3.8, 1.0, 0.28, size=13, color=CYAN, bold=True)
    add_text(
        s,
        "cluster history:  2s, 5s, 40s, 90s",
        1.42,
        4.26,
        4.5,
        0.34,
        size=15,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    add_text(
        s,
        "elapsed = 10s  → survivors = 40, 90",
        1.42,
        4.73,
        4.5,
        0.3,
        size=14.5,
        color=MUTED,
        font=MONO,
    )
    add_text(
        s,
        "prediction = mean(30, 80) = 55s remaining",
        1.42,
        5.15,
        4.65,
        0.32,
        size=14.5,
        color=GREEN,
        bold=True,
        font=MONO,
    )
    rect(s, 6.92, 3.52, 5.3, 2.25, fill=PANEL, line=CORAL)
    badge(s, "BOUNDARY", 7.22, 3.8, 1.18, fill=CORAL, color=BG)
    add_text(
        s,
        "论文的 tool_args = 完整外层 payload",
        7.2,
        4.35,
        4.5,
        0.38,
        size=19,
        bold=True,
    )
    add_text(
        s,
        "不解析 shell AST · 不拆 pipeline · 不观察 execve child · 不使用 per-process counters",
        7.22,
        4.95,
        4.55,
        0.62,
        size=15,
        color=MUTED,
    )
    footer(
        s,
        "CacheWise paper §5.2 and footnotes 2–3 · analysis/offline/related-work.md, “CacheWise granularity”.",
    )

    # 6 — claims and scope
    s = new_slide(prs)
    title(
        s,
        "论文的 3.5× headline ≠ argument predictor 的 isolated gain",
        6,
        kicker="Claim boundary",
    )
    rect(s, 0.7, 1.43, 6.0, 4.98, fill=PANEL, line=GREEN)
    badge(s, "PAPER · FULL SYSTEM", 1.02, 1.7, 1.95, fill=GREEN, color=BG)
    add_text(
        s,
        "2× H200 · Qwen2.5-Coder-32B · TP=2",
        1.02,
        2.28,
        5.2,
        0.34,
        size=17,
        color=WHITE,
        bold=True,
    )
    bullet_list(
        s,
        [
            "CATraces random 80/20 sessions",
            "30 / 40 / 50 concurrent sessions",
            "prefix-aware scheduling + predictive eviction",
        ],
        1.05,
        2.86,
        5.1,
        1.65,
        size=15,
    )
    stat_card(
        s, 1.02, 4.75, 1.65, "2.7–3.5×", "session time", color=GREEN, value_size=18
    )
    stat_card(s, 2.9, 4.75, 1.55, "2–2.6×", "fewer evictions", color=GREEN)
    stat_card(s, 4.68, 4.75, 1.55, "≤19%", "C100 isolated", color=YELLOW)
    rect(s, 6.92, 1.43, 5.72, 4.98, fill=PANEL, line=CORAL)
    badge(s, "OURS · PREDICTOR TRANSFER", 7.24, 1.7, 2.48, fill=CORAL, color=BG)
    add_text(
        s,
        "只复现一个必要子命题",
        7.22,
        2.28,
        4.65,
        0.35,
        size=18,
        color=WHITE,
        bold=True,
    )
    add_text(
        s,
        "“whole args 是否比 tool-name 更会排 victim？”",
        7.22,
        2.75,
        4.72,
        0.55,
        size=21,
        color=CORAL,
        bold=True,
    )
    bullet_list(
        s,
        [
            "同样的 global / tool / C20 / C50 / C100",
            "同样的 conditional remaining 与 C100 primary",
            "没有实际 KV capacity / eviction / JCT",
        ],
        7.22,
        3.52,
        4.8,
        1.72,
        size=15,
    )
    add_text(
        s,
        "如果 predictor gate 不过，就没有理由花成本做 live serving。",
        7.22,
        5.55,
        4.85,
        0.5,
        size=16,
        color=YELLOW,
        bold=True,
    )
    footer(
        s,
        "CacheWise paper §6, Figures 13–19. Paper leaves temporal/cross-project drift robustness to future work.",
    )

    # 7 — datasets
    s = new_slide(prs)
    title(
        s, "我们的 traces 保留 raw args，但物理条件与论文不同", 7, kicker="Data & setup"
    )
    headers = ["corpus", "fit", "evaluation", "tool gaps", "observed ranking"]
    xs = [0.78, 3.3, 5.05, 7.25, 9.3]
    widths = [2.35, 1.55, 2.0, 1.85, 3.05]
    for x, w, text_ in zip(xs, widths, headers, strict=True):
        rect(s, x, 1.56, w, 0.55, fill=PANEL2, line=PANEL2, radius=False)
        add_text(
            s,
            text_,
            x + 0.08,
            1.71,
            w - 0.16,
            0.26,
            size=12,
            color=CYAN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    rows = [
        (
            "SWE-ReBench",
            "100 sessions",
            "277 sessions",
            "4,175 → 12,771",
            "2,410 events · max c=2",
        ),
        (
            "Terminal-Bench",
            "139 declared\n136 w/ gaps",
            "100 declared\n99 w/ gaps",
            "1,888 → 1,227",
            "synthetic only",
        ),
    ]
    for i, row in enumerate(rows):
        yy = 2.22 + i * 0.92
        for j, (x, w, text_) in enumerate(zip(xs, widths, row, strict=True)):
            rect(
                s,
                x,
                yy,
                w,
                0.78,
                fill=PANEL if i == 0 else "10192B",
                line=PANEL2,
                radius=False,
            )
            add_text(
                s,
                text_,
                x + 0.08,
                yy + 0.14,
                w - 0.16,
                0.48,
                size=13 if j else 14,
                color=WHITE,
                bold=j == 0,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
    rect(s, 0.78, 4.28, 3.55, 1.82, fill=PANEL, line=GREEN)
    badge(s, "WE HAVE", 1.05, 4.55, 1.0, fill=GREEN, color=BG)
    add_text(
        s,
        "raw outer tool args\n≈0.5s causal CPU/network samples",
        1.05,
        5.02,
        2.95,
        0.82,
        size=14.5,
        color=WHITE,
        bold=True,
    )
    rect(s, 4.6, 4.28, 3.55, 1.82, fill=PANEL, line=CORAL)
    badge(s, "WE DO NOT HAVE", 4.88, 4.55, 1.55, fill=CORAL, color=BG)
    add_text(
        s,
        "recorded shared KV pressure\nactual eviction / recompute feedback",
        4.88,
        5.02,
        2.95,
        0.82,
        size=14.5,
        color=WHITE,
        bold=True,
    )
    rect(s, 8.43, 4.28, 3.92, 1.82, fill=PANEL, line=YELLOW)
    badge(s, "SYNTHETIC", 8.72, 4.55, 1.25, fill=YELLOW, color=BG)
    add_text(
        s,
        "只平移真实 trace 的到达时间\n内部 gap duration 保持不变",
        8.72,
        5.02,
        3.15,
        0.82,
        size=14.5,
        color=WHITE,
        bold=True,
    )
    footer(
        s,
        "Local artifacts: SWE/TB result.json data fields; analysis/offline/related-work.md. All corpora are development-exposed.",
    )

    # 8 — metric
    s = new_slide(prs)
    title(
        s,
        "唯一主指标：选错 victim 会损失多少“可释放时间”",
        8,
        kicker="Metric · lower is better",
    )
    add_text(
        s, "ranking event @ t", 0.75, 1.48, 2.0, 0.3, size=13, color=CYAN, bold=True
    )
    candidates = [("A", 40, 10, BLUE), ("B", 15, 20, YELLOW), ("C", 4, 5, PURPLE)]
    for i, (name, actual, pred, color) in enumerate(candidates):
        yy = 2.05 + i * 0.88
        rect(s, 0.82, yy, 5.2, 0.65, fill=PANEL, line=color)
        add_text(s, name, 1.05, yy + 0.17, 0.42, 0.25, size=16, color=color, bold=True)
        add_text(
            s,
            f"actual remaining  {actual:>2}s",
            1.65,
            yy + 0.14,
            1.95,
            0.28,
            size=15,
            color=WHITE,
            font=MONO,
        )
        add_text(
            s,
            f"tool predicts  {pred:>2}s",
            3.72,
            yy + 0.14,
            1.88,
            0.28,
            size=15,
            color=MUTED,
            font=MONO,
        )
    add_text(
        s,
        "tool chooses B",
        1.35,
        5.15,
        1.75,
        0.32,
        size=17,
        color=YELLOW,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "→",
        3.1,
        5.12,
        0.6,
        0.34,
        size=20,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "regret = 40 − 15 = 25s",
        3.62,
        5.14,
        2.72,
        0.34,
        size=17,
        color=RED,
        bold=True,
    )
    rect(s, 6.62, 1.56, 5.66, 4.66, fill=PANEL, line=PANEL2)
    add_text(s, "定义", 6.98, 1.9, 0.8, 0.3, size=13, color=CYAN, bold=True)
    add_text(
        s,
        "regret = max(actual remaining)\n              − chosen actual remaining",
        7.0,
        2.3,
        4.8,
        1.08,
        size=17.5,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    add_text(
        s, "0 = oracle choice", 7.0, 3.58, 2.6, 0.3, size=15, color=GREEN, bold=True
    )
    add_text(
        s,
        "Δ = mean regret(candidate) − mean regret(tool)",
        7.0,
        4.08,
        4.78,
        0.42,
        size=17,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    badge(s, "GO only if 95% CI upper < 0", 7.0, 4.78, 3.35, fill=GREEN, color=BG)
    add_text(
        s,
        "口径：hypothetical victim ordering\n≠ eviction count  ≠ KV bytes  ≠ JCT",
        7.0,
        5.34,
        4.78,
        0.66,
        size=15,
        color=MUTED,
    )
    footer(
        s,
        "Local protocol: analysis/offline/related-work.md; src/tool_resource_eval/cachewise_reproduction.py::evaluate.",
    )

    # 9 — experimental ladder
    s = new_slide(prs)
    title(s, "五个冻结实验逐步排除四种解释", 9, kicker="Decision ladder")
    experiments = [
        ("1", "SWE observed overlap", "whole args?", "+0.587s", "NO-GO", RED),
        ("2", "causal progress", "live counters?", "+0.505s", "NO-GO", RED),
        ("3", "SWE synthetic c32", "overlap too sparse?", "+1.847s", "NO-GO", RED),
        ("4", "TB synthetic c32", "corpus-specific?", "−72.299s", "FRAGILE", YELLOW),
        (
            "5",
            "exact timeout + LOTO",
            "simple signal stable?",
            "+4.012s*",
            "NO-GO",
            RED,
        ),
    ]
    for i, (num, name, question, delta, verdict, color) in enumerate(experiments):
        yy = 1.42 + i * 1.02
        circle(s, 0.82, yy + 0.12, 0.46, fill=color)
        add_text(
            s,
            num,
            0.82,
            yy + 0.19,
            0.46,
            0.2,
            size=11,
            color=BG,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if i < len(experiments) - 1:
            line(s, 1.05, yy + 0.58, 1.05, yy + 1.1, color=PANEL2, width=3)
        rect(s, 1.55, yy, 10.78, 0.78, fill=PANEL, line=color)
        add_text(s, name, 1.8, yy + 0.18, 2.55, 0.3, size=17, color=WHITE, bold=True)
        add_text(s, question, 4.45, yy + 0.19, 2.65, 0.28, size=14, color=MUTED)
        add_text(
            s,
            delta,
            7.45,
            yy + 0.16,
            1.55,
            0.3,
            size=17,
            color=color,
            bold=True,
            align=PP_ALIGN.RIGHT,
            font=MONO,
        )
        badge(s, verdict, 9.55, yy + 0.18, 1.55, fill=color, color=BG)
    add_text(
        s,
        "* step 5 reports the predeclared worst delete-one-task result; aggregate timeout Δ = −17.271s.",
        1.55,
        6.67,
        9.9,
        0.25,
        size=10,
        color=FAINT,
    )
    footer(
        s,
        "Frozen protocols and outcomes: analysis/development/research-directions-20260729.md.",
    )

    # 10 — SWE reproduction
    s = new_slide(prs)
    title(
        s,
        "SWE：argument clustering 增大 regret，主 gate 失败",
        10,
        kicker="Result 1 · observed overlap",
    )
    swe_metrics = SWE["metrics"]
    labels = ["global", "tool", "C20", "C50", "C100"]
    values = [
        swe_metrics[k]["mean_regret_s"]
        for k in ["global", "tool", "c20", "c50", "c100"]
    ]
    hbars(
        s,
        labels,
        values,
        0.82,
        1.55,
        6.15,
        3.72,
        colors=[FAINT, BLUE, CYAN, YELLOW, CORAL],
    )
    p = SWE["primary_comparison"]
    rect(s, 7.38, 1.55, 4.88, 1.55, fill=PANEL, line=RED)
    add_text(
        s, "PRIMARY  C100 − tool", 7.7, 1.82, 3.8, 0.28, size=12, color=MUTED, bold=True
    )
    add_text(
        s,
        f"+{p['delta_mean_regret_s']:.3f} s",
        7.68,
        2.18,
        2.3,
        0.46,
        size=29,
        color=RED,
        bold=True,
    )
    add_text(
        s,
        "95% CI\n[+0.120, +1.188]",
        10.02,
        2.09,
        1.82,
        0.7,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ci_axis(
        s,
        7.7,
        3.72,
        4.12,
        0.120,
        0.587,
        1.188,
        min_v=-0.3,
        max_v=1.3,
        label="Δ mean victim regret (s)",
        color=RED,
    )
    rect(s, 7.38, 4.52, 4.88, 1.43, fill=PANEL, line=PANEL2)
    add_text(
        s,
        "rare tail dominates",
        7.68,
        4.76,
        2.2,
        0.28,
        size=13,
        color=YELLOW,
        bold=True,
    )
    add_text(
        s,
        "top-1: 93.98% → 94.07%",
        7.68,
        5.18,
        2.28,
        0.28,
        size=14,
        color=WHITE,
        bold=True,
    )
    add_text(
        s,
        "p99 regret: 17.3s → 35.5s",
        9.82,
        5.18,
        2.15,
        0.28,
        size=14,
        color=RED,
        bold=True,
    )
    add_text(
        s,
        "90 choice changes: 46 help / 44 harm · harm 1,803s vs help 390s",
        0.86,
        6.25,
        11.2,
        0.34,
        size=15,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    footer(
        s,
        "Source: analysis/results/cachewise-swe-reproduction-20260731/result.json; tail diagnostic in analysis/offline/related-work.md.",
    )

    # 11 — SWE failure mechanism
    s = new_slide(prs)
    title(
        s,
        "SWE 失败不是“cluster 太小”，而是相似命令跨 repo 的尺度不同",
        11,
        kicker="Failure analysis · evidence vs inference",
    )
    rect(s, 0.72, 1.45, 6.15, 4.98, fill=PANEL, line=CORAL)
    badge(s, "OBSERVED", 1.02, 1.72, 1.2, fill=CORAL, color=BG)
    add_text(
        s,
        "syntactically similar pytest / test-suite cluster",
        1.02,
        2.24,
        5.25,
        0.38,
        size=17,
        color=WHITE,
        bold=True,
    )
    add_text(s, "FIT cluster max", 1.06, 3.03, 1.7, 0.3, size=13, color=MUTED)
    add_text(s, "3–29s", 2.76, 2.91, 1.6, 0.46, size=27, color=YELLOW, bold=True)
    line(s, 1.08, 3.55, 5.98, 3.55, color=PANEL2, width=1)
    add_text(s, "HELD-OUT still active", 1.06, 3.84, 2.1, 0.3, size=13, color=MUTED)
    add_text(s, "108–300s", 3.2, 3.72, 2.25, 0.46, size=27, color=CORAL, bold=True)
    add_text(
        s,
        "C100 predicts 0–5s remaining",
        1.06,
        4.63,
        3.05,
        0.34,
        size=16,
        color=RED,
        bold=True,
    )
    add_text(
        s,
        "tool-name predicts 20–53s",
        3.55,
        4.63,
        2.65,
        0.34,
        size=16,
        color=GREEN,
        bold=True,
    )
    add_text(
        s,
        "median harmful-cluster surviving support = 36\n→ 不是简单 singleton 问题",
        1.06,
        5.25,
        5.2,
        0.63,
        size=15,
        color=MUTED,
    )
    rect(s, 7.12, 1.45, 5.5, 2.18, fill=PANEL, line=YELLOW)
    badge(s, "SUPPORTED INFERENCE", 7.43, 1.72, 1.85, fill=YELLOW, color=BG)
    add_text(
        s,
        "TF-IDF 识别“在跑测试”，\n却看不到 suite scale、dependency/cache state、repo-specific work。",
        7.42,
        2.28,
        4.6,
        0.9,
        size=18,
        color=WHITE,
        bold=True,
    )
    rect(s, 7.12, 3.9, 5.5, 2.53, fill=PANEL, line=FAINT)
    badge(s, "NOT TESTED", 7.43, 4.2, 1.28, fill=FAINT, color=BG)
    add_text(
        s,
        "环境状态特征是否能泛化？",
        7.42,
        4.78,
        4.4,
        0.35,
        size=18,
        color=WHITE,
        bold=True,
    )
    add_text(
        s,
        "repo size · test inventory · dependency/cache state · process phase",
        7.42,
        5.32,
        4.55,
        0.65,
        size=15,
        color=MUTED,
    )
    add_text(
        s,
        "不能在已暴露 fresh-277 上继续调 tokenization / support threshold。",
        7.42,
        5.82,
        4.55,
        0.48,
        size=11.5,
        color=RED,
        bold=True,
    )
    footer(
        s,
        "Source: analysis/offline/related-work.md, post-result mechanism diagnostic. Top-10 task pairs explain 91.5% of harmful delta.",
    )

    # 12 — progress
    s = new_slide(prs)
    title(
        s,
        "Progress counters 改善 MAE，却把 ranking 决策推向错误方向",
        12,
        kicker="Result 2 · causal progress",
    )
    pm = PROGRESS["metrics"]
    hbars(
        s,
        ["tool", "progress"],
        [pm["tool"]["mean_regret_s"], pm["progress"]["mean_regret_s"]],
        0.82,
        1.55,
        5.85,
        1.85,
        colors=[BLUE, CORAL],
    )
    ci_axis(
        s,
        0.95,
        4.15,
        5.2,
        0.257,
        0.505,
        0.795,
        min_v=-0.2,
        max_v=0.9,
        label="progress − tool regret (s)",
        color=RED,
    )
    rect(s, 6.95, 1.52, 5.38, 4.85, fill=PANEL, line=PANEL2)
    add_text(
        s, "point prediction", 7.3, 1.82, 1.75, 0.28, size=13, color=CYAN, bold=True
    )
    stat_card(
        s,
        7.28,
        2.28,
        2.15,
        "57.43→55.24",
        "MAE improves",
        color=GREEN,
        value_size=18,
    )
    stat_card(
        s,
        9.72,
        2.28,
        2.15,
        "+26.25→−16.83",
        "signed error flips",
        color=CORAL,
        value_size=14.5,
    )
    add_text(
        s, "decision boundary", 7.3, 3.68, 1.75, 0.28, size=13, color=YELLOW, bold=True
    )
    stat_card(s, 7.28, 4.13, 2.15, "8 / 63", "help / harm changes", color=RED)
    stat_card(
        s,
        9.72,
        4.13,
        2.15,
        "32.9 / 1,249",
        "help / harm seconds",
        color=RED,
        value_size=18,
    )
    add_text(
        s,
        "OBSERVED: 71 changes都从 still-running gap 转向新到达 gap。",
        7.3,
        5.62,
        4.55,
        0.34,
        size=14,
        color=WHITE,
        bold=True,
    )
    add_text(
        s,
        "INFERENCE: model 对活跃长任务变得过度乐观；fit metric 与 policy utility 方向相反。",
        7.3,
        6.05,
        4.55,
        0.42,
        size=14,
        color=YELLOW,
    )
    footer(
        s,
        "Source: analysis/results/cachewise-progress-counter-20260731/result.json. Features are causal ≈0.5s CPU/network samples.",
    )

    # 13 — synthetic SWE
    s = new_slide(prs)
    title(
        s,
        "把 concurrency 合成到 c32，SWE 仍然失败",
        13,
        kicker="Result 3 · fixed-duration counterfactual",
    )
    for i, (name, start, length, color) in enumerate(
        [
            ("task A", 0.0, 4.8, BLUE),
            ("task B", 1.0, 3.4, YELLOW),
            ("task C", 2.0, 4.3, PURPLE),
            ("task D", 3.0, 2.5, CYAN),
        ]
    ):
        yy = 1.65 + i * 0.48
        add_text(s, name, 0.78, yy - 0.05, 0.75, 0.25, size=10.5, color=MUTED)
        rect(
            s,
            1.55 + start * 0.63,
            yy,
            length * 0.63,
            0.17,
            fill=color,
            line=color,
            radius=False,
        )
    add_text(
        s,
        "arrival shifts only; each trace keeps its internal gaps + durations",
        0.82,
        3.72,
        5.3,
        0.32,
        size=13,
        color=YELLOW,
        bold=True,
    )
    d = SWE_SYNTH["data"]
    stat_card(s, 0.82, 4.2, 1.72, "32", "arrival seeds", color=CYAN)
    stat_card(
        s,
        2.76,
        4.2,
        1.72,
        f"{d['realized_mean_live_sessions_mean']:.1f}",
        "mean live sessions",
        color=CYAN,
    )
    stat_card(s, 4.7, 4.2, 1.72, "36–42", "max live sessions", color=CYAN)
    rect(s, 6.82, 1.52, 5.55, 4.88, fill=PANEL, line=RED)
    add_text(
        s, "SWE synthetic-c32", 7.17, 1.86, 2.6, 0.35, size=18, color=WHITE, bold=True
    )
    add_text(s, "tool", 7.18, 2.58, 1.0, 0.25, size=13, color=MUTED)
    add_text(s, "57.961s", 8.05, 2.47, 1.5, 0.38, size=22, color=BLUE, bold=True)
    add_text(s, "C100", 9.73, 2.58, 1.0, 0.25, size=13, color=MUTED)
    add_text(s, "59.808s", 10.62, 2.47, 1.5, 0.38, size=22, color=CORAL, bold=True)
    p = SWE_SYNTH["primary_comparison"]
    add_text(
        s,
        f"Δ +{p['delta_mean_regret_s']:.3f}s",
        7.18,
        3.42,
        2.2,
        0.42,
        size=25,
        color=RED,
        bold=True,
    )
    add_text(
        s,
        "95% CI [+1.087, +2.646]",
        9.3,
        3.5,
        2.55,
        0.3,
        size=14,
        color=WHITE,
        bold=True,
    )
    badge(s, "24 harmful / 8 helpful schedules", 7.18, 4.3, 3.58, fill=RED, color=BG)
    add_text(
        s,
        "已排除：natural SWE overlap 太稀疏",
        7.18,
        5.0,
        4.25,
        0.4,
        size=14,
        color=WHITE,
        bold=True,
    )
    add_text(
        s,
        "未排除：真实 contention 会改变 duration / arrival / eviction feedback。",
        7.18,
        5.48,
        4.45,
        0.55,
        size=13,
        color=MUTED,
    )
    footer(
        s,
        "Source: analysis/results/cachewise-swe-synthetic-c32-20260731/result.json. This is not a live concurrency experiment.",
    )

    # 14 — TB C100 fragility
    s = new_slide(prs)
    title(
        s,
        "TB aggregate 反转为正，但几乎由 hdfs-deployment 扛住",
        14,
        kicker="Result 4 · cross-corpus transfer",
    )
    tm = TB_SYNTH["schedule_mean_metrics"]
    hbars(
        s,
        ["tool", "C100"],
        [tm["tool"]["mean_regret_s"], tm["c100"]["mean_regret_s"]],
        0.75,
        1.52,
        5.7,
        1.9,
        colors=[BLUE, GREEN],
    )
    p = TB_SYNTH["primary_comparison"]
    ci_axis(
        s,
        0.95,
        4.08,
        5.15,
        -90.844,
        -72.299,
        -53.768,
        min_v=-105,
        max_v=45,
        label="C100 − tool regret (s)",
        color=GREEN,
    )
    add_text(
        s,
        "29 / 32 schedules improve",
        1.55,
        5.03,
        3.85,
        0.32,
        size=17,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    rect(s, 6.72, 1.48, 5.72, 2.2, fill=PANEL, line=CORAL)
    badge(s, "POST-HOC SENSITIVITY", 7.03, 1.75, 2.1, fill=CORAL, color=BG)
    add_text(
        s,
        "remove hdfs-deployment",
        7.03,
        2.3,
        2.65,
        0.3,
        size=16,
        color=MUTED,
        bold=True,
    )
    add_text(s, "+26.722 s", 9.72, 2.17, 2.0, 0.48, size=28, color=RED, bold=True)
    add_text(
        s,
        "95% CI [+17.355, +37.789] · harmful 28/32",
        7.03,
        2.91,
        4.75,
        0.35,
        size=14,
        color=WHITE,
        bold=True,
    )
    rect(s, 6.72, 3.95, 5.72, 2.42, fill=PANEL, line=YELLOW)
    badge(s, "OBSERVED", 7.03, 4.22, 1.2, fill=YELLOW, color=BG)
    add_text(
        s,
        "hdfs-deployment",
        7.03,
        4.75,
        2.5,
        0.32,
        size=19,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    add_text(
        s,
        "2 × ≈600s download exec",
        9.42,
        4.73,
        2.3,
        0.34,
        size=18,
        color=YELLOW,
        bold=True,
    )
    add_text(
        s,
        "C100 从 command text + explicit timeout=600 识别长 survivor。",
        7.03,
        5.3,
        4.9,
        0.42,
        size=15.5,
        color=MUTED,
    )
    add_text(
        s,
        f"C50 regret {tm['c50']['mean_regret_s']:.3f}s < C100 {tm['c100']['mean_regret_s']:.3f}s；但 C50 不是 primary，不能事后改选。",
        7.03,
        5.82,
        4.92,
        0.42,
        size=13,
        color=CORAL,
        bold=True,
    )
    footer(
        s,
        "Source: analysis/results/cachewise-tb-synthetic-c32-20260731/result.json; hdfs sensitivity in research-directions-20260729.md.",
    )

    # 15 — timeout stability
    s = new_slide(prs)
    title(
        s,
        "Timeout 只是 ceiling：task-stability gate 失败",
        15,
        kicker="Result 5 · exact-timeout LOTO",
    )
    tp = TIMEOUT["primary_comparison"]
    rect(s, 0.72, 1.48, 5.85, 2.05, fill=PANEL, line=GREEN)
    badge(s, "AGGREGATE PASSES", 1.02, 1.77, 1.7, fill=GREEN, color=BG)
    add_text(s, "tool 194.633s", 1.02, 2.36, 1.9, 0.32, size=16, color=BLUE, bold=True)
    add_text(
        s,
        "→",
        2.83,
        2.34,
        0.5,
        0.32,
        size=18,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s, "timeout 177.362s", 3.28, 2.36, 2.2, 0.32, size=16, color=GREEN, bold=True
    )
    add_text(
        s,
        f"Δ {tp['delta_mean_regret_s']:.3f}s · CI [−29.739, −4.296] · 20/32 helpful",
        1.02,
        2.9,
        4.9,
        0.32,
        size=15,
        color=WHITE,
        bold=True,
    )
    rect(s, 0.72, 3.82, 5.85, 2.3, fill=PANEL, line=RED)
    badge(s, "TASK STABILITY FAILS", 1.02, 4.1, 1.85, fill=RED, color=BG)
    add_text(
        s,
        "delete mixed-integer-programming",
        1.02,
        4.7,
        3.5,
        0.32,
        size=17,
        color=WHITE,
        bold=True,
        font=MONO,
    )
    add_text(s, "+4.012s", 4.55, 4.58, 1.55, 0.45, size=24, color=RED, bold=True)
    add_text(
        s,
        "仅 13/32 schedules helpful · 99 个 task 全穷举，只有这一个翻转符号",
        1.02,
        5.33,
        4.95,
        0.48,
        size=14.5,
        color=MUTED,
    )
    rect(s, 6.88, 1.48, 5.55, 4.64, fill=PANEL, line=YELLOW)
    add_text(
        s,
        "timeout = 300s",
        7.25,
        1.9,
        2.35,
        0.38,
        size=22,
        color=YELLOW,
        bold=True,
        font=MONO,
    )
    line(s, 7.3, 2.78, 11.88, 2.78, color=FAINT, width=3)
    line(s, 11.2, 2.47, 11.2, 3.12, color=YELLOW, width=3)
    add_text(
        s,
        "ceiling",
        10.65,
        3.18,
        1.1,
        0.26,
        size=12,
        color=YELLOW,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    for x, d, label, color in [
        (7.55, 0.13, "0.1s", BLUE),
        (8.35, 0.18, "4s", CYAN),
        (9.45, 0.25, "90s", PURPLE),
        (11.08, 0.28, "300.011s", CORAL),
    ]:
        circle(s, x, 2.64, d, fill=color)
        add_text(
            s,
            label,
            x - 0.28,
            3.55,
            0.95,
            0.24,
            size=10.5,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        s,
        "同一个 timeout 值覆盖多个数量级的实际 work。",
        7.25,
        4.16,
        4.65,
        0.34,
        size=17,
        color=WHITE,
        bold=True,
    )
    badge(s, "OBSERVED", 7.25, 4.78, 1.2, fill=CYAN, color=BG)
    add_text(
        s,
        "600s signature 提供最大 aggregate benefit；\n300s signature 若没有这个跑满 300s 的 solver call，则净有害。",
        7.25,
        5.3,
        4.62,
        0.64,
        size=15,
        color=MUTED,
    )
    footer(
        s,
        "Source: analysis/results/cachewise-tb-timeout-stability-20260731/result.json; post-hoc attribution in research-directions-20260729.md.",
    )

    # 16 — conclusions and next steps
    s = new_slide(prs)
    title(
        s,
        "结论：停止 tuning；下一次证据必须跨过 task 与真实压力两道门",
        16,
        kicker="Implications · takeaways · next steps",
    )
    takeaways = [
        (
            "① insight survives",
            "相对 next-reuse ordering 是正确 systems objective；tool metadata 确实含 tail signal。",
            GREEN,
        ),
        (
            "② estimator fails",
            "whole-args、progress、synthetic SWE、exact timeout 都未通过稳定 gate。",
            RED,
        ),
        (
            "③ scope matters",
            "offline regret 不是 eviction/JCT；TB aggregate pass 也不能覆盖 single-task fragility。",
            YELLOW,
        ),
    ]
    for i, (head, body, color) in enumerate(takeaways):
        x = 0.72 + i * 4.18
        rect(s, x, 1.48, 3.88, 1.82, fill=PANEL, line=color)
        add_text(s, head, x + 0.22, 1.75, 3.35, 0.3, size=14, color=color, bold=True)
        add_text(s, body, x + 0.22, 2.18, 3.42, 0.82, size=13.5, color=WHITE, bold=True)
    add_text(
        s, "NEXT EVIDENCE RUNG", 0.75, 3.62, 2.25, 0.28, size=12, color=CYAN, bold=True
    )
    stages = [
        ("freeze current NO-GO", "no post-hoc clause/text rescue", RED),
        (
            "new independent tasks",
            "task-grouped gate; environment-scale hypothesis",
            BLUE,
        ),
        ("task-stable ranker", "charge inference + scheduling overhead", YELLOW),
        ("live KV pressure", "30–50 sessions; actual evictions + JCT", GREEN),
    ]
    for i, (head, body, color) in enumerate(stages):
        x = 0.72 + i * 3.08
        rect(s, x, 4.08, 2.72, 1.25, fill=PANEL, line=color)
        add_text(
            s,
            head,
            x + 0.18,
            4.32,
            2.34,
            0.3,
            size=14,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            s,
            body,
            x + 0.16,
            4.73,
            2.4,
            0.4,
            size=11.5,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
        if i < len(stages) - 1:
            add_text(
                s,
                "→",
                x + 2.73,
                4.48,
                0.35,
                0.3,
                size=18,
                color=FAINT,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    rect(s, 0.72, 5.75, 11.95, 0.68, fill=PANEL2, line=PANEL2)
    add_text(s, "最终 takeaway", 0.98, 5.96, 1.35, 0.26, size=12, color=CYAN, bold=True)
    add_text(
        s,
        "“use tool args” 不是方法贡献；可发表的问题是：什么 inference-time state 能跨 task 稳定改善一个真实、计费完整的 KV decision？",
        2.42,
        5.88,
        9.75,
        0.42,
        size=16,
        color=WHITE,
        bold=True,
    )
    footer(
        s,
        "Paper: arXiv:2606.16824v1 · Local source index: analysis/offline/related-work.md and analysis/development/research-directions-20260729.md.",
    )

    return prs


def validate_bounds(prs: Presentation) -> None:
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            x = shape.left / Inches(1)
            y = shape.top / Inches(1)
            w = shape.width / Inches(1)
            h = shape.height / Inches(1)
            if x < -1e-6 or y < -1e-6 or x + w > W + 1e-6 or y + h > H + 1e-6:
                raise ValueError(
                    f"slide {slide_index} shape {shape.name!r} out of bounds: "
                    f"({x:.2f}, {y:.2f}, {w:.2f}, {h:.2f})"
                )


if __name__ == "__main__":
    presentation = generate()
    validate_bounds(presentation)
    presentation.save(OUT)
    print(f"wrote {OUT} ({len(presentation.slides)} slides)")
