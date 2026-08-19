#!/usr/bin/env python3
"""Generate the editable Chinese research-journey deck.

Run from the repository root:
  uv run --with python-pptx python analysis/slides/research-journey-20260819/generate_slides.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUT = HERE / "agent-resource-research-journey-zh-20260819.pptx"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
MONO = "Liberation Mono"

BG = "09111F"
PANEL = "111D30"
PANEL2 = "17263D"
WHITE = "F4F7FB"
MUTED = "A7B6CA"
FAINT = "677990"
CYAN = "32D5C4"
BLUE = "62A8FF"
GREEN = "40D39C"
YELLOW = "F4C75B"
CORAL = "FF7A85"
PURPLE = "B69BFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = WHITE,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = FONT,
    margin: float = 0.03,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    return box


def box(slide, x, y, w, h, *, fill=PANEL, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def rule(slide, x1, y1, x2, y2, *, color=FAINT, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def arrow(slide, x, y, w, h, *, fill=CYAN):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(fill)
    return shape


def circle(slide, x, y, d, *, fill=CYAN):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(fill)
    return shape


def add_title(slide, title: str, page: int, kicker: str):
    text(slide, kicker, 0.62, 0.22, 7.0, 0.24, size=9.5, color=CYAN, bold=True)
    text(slide, title, 0.60, 0.48, 12.0, 0.66, size=26, bold=True)
    rule(slide, 0.62, 1.15, 12.72, 1.15, color=PANEL2, width=1)
    text(
        slide,
        f"{page:02d}",
        12.38,
        0.22,
        0.35,
        0.24,
        size=9.5,
        color=FAINT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_source(slide, value: str):
    rule(slide, 0.62, 7.06, 12.72, 7.06, color=PANEL2, width=0.8)
    text(slide, value, 0.62, 7.10, 12.0, 0.20, size=7.8, color=FAINT)


def add_bullets(slide, items, x, y, w, h, *, size=16, color=WHITE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(9)
    return shape


def chip(slide, value, x, y, w, *, fill=PANEL2, color=WHITE):
    box(slide, x, y, w, 0.34, fill=fill, line=fill)
    text(
        slide,
        value,
        x + 0.05,
        y + 0.02,
        w - 0.1,
        0.25,
        size=10.2,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def stat(slide, x, y, w, value, label, *, color=CYAN, note=""):
    box(slide, x, y, w, 1.12, fill=PANEL, line=PANEL2)
    text(slide, value, x + 0.15, y + 0.12, w - 0.3, 0.38, size=23, color=color, bold=True)
    text(slide, label, x + 0.15, y + 0.54, w - 0.3, 0.24, size=11.5, bold=True)
    if note:
        text(slide, note, x + 0.15, y + 0.80, w - 0.3, 0.20, size=8.8, color=MUTED)


def lesson(slide, value: str, *, color=YELLOW):
    box(slide, 0.72, 6.27, 11.90, 0.58, fill=PANEL2, line=color)
    text(slide, "这一阶段学到", 0.92, 6.43, 1.45, 0.22, size=10.5, color=color, bold=True)
    text(slide, value, 2.26, 6.38, 10.0, 0.30, size=14.2, bold=True)


def bar_chart(slide, labels, values, x, y, w, h, *, colors=None, suffix="%", max_value=None):
    colors = colors or [CYAN] * len(values)
    maximum = max_value or max(values) * 1.10
    row_h = h / len(values)
    label_w = 1.75
    bar_w = w - label_w - 0.85
    for i, (label, value, color) in enumerate(zip(labels, values, colors, strict=True)):
        yy = y + i * row_h
        text(slide, label, x, yy + 0.04, label_w - 0.1, 0.28, size=12.3, color=MUTED)
        box(slide, x + label_w, yy + 0.07, bar_w, 0.25, fill=PANEL2, line=PANEL2, radius=False)
        box(
            slide,
            x + label_w,
            yy + 0.07,
            max(0.04, bar_w * value / maximum),
            0.25,
            fill=color,
            line=color,
            radius=False,
        )
        text(
            slide,
            f"{value:,.3f}{suffix}" if isinstance(value, float) else f"{value:,}{suffix}",
            x + label_w + bar_w + 0.12,
            yy + 0.02,
            0.72,
            0.32,
            size=11.2,
            color=color,
            bold=True,
        )


def card(slide, x, y, w, h, heading, body, *, color=CYAN, heading_size=15, body_size=12.5):
    box(slide, x, y, w, h, fill=PANEL, line=PANEL2)
    rule(slide, x + 0.15, y + 0.18, x + 0.15, y + h - 0.18, color=color, width=3.5)
    text(slide, heading, x + 0.34, y + 0.16, w - 0.5, 0.42, size=heading_size, color=color, bold=True)
    text(slide, body, x + 0.34, y + 0.65, w - 0.5, h - 0.78, size=body_size, color=MUTED)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(BG)
    return slide


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # 1 — title
    s = new_slide(prs)
    chip(s, "组会版 · 15–20 分钟", 0.70, 0.58, 1.85, fill=CYAN, color=BG)
    text(s, "从“猜工具多久”\n到“让 CPU 与 GPU 协同”", 0.68, 1.22, 8.8, 1.55, size=31, bold=True)
    text(
        s,
        "一条由失败结果推动的研究路线：预测什么、为什么失败、怎样变成真正有用的动作",
        0.72,
        2.92,
        9.1,
        0.65,
        size=17,
        color=MUTED,
    )
    phases = [("文本相似度", CORAL), ("真实工作量", YELLOW), ("资源预测", BLUE), ("调度动作", GREEN)]
    for i, (label, color) in enumerate(phases):
        x = 0.82 + i * 2.95
        circle(s, x, 5.06, 0.32, fill=color)
        text(s, label, x - 0.26, 5.52, 1.65, 0.34, size=13, color=color, bold=True)
        if i < len(phases) - 1:
            arrow(s, x + 0.64, 5.12, 1.72, 0.20, fill=PANEL2)
    text(s, "2026-08-19", 10.75, 6.60, 1.75, 0.28, size=11, color=FAINT, align=PP_ALIGN.RIGHT)

    # 2 — mental model
    s = new_slide(prs)
    add_title(s, "Agent 不是一直在用 GPU：它在“思考”和“做事”之间切换", 2, "先建立共同直觉")
    text(s, "一个任务的真实时间线", 0.75, 1.45, 2.8, 0.35, size=15, color=MUTED, bold=True)
    labels = [
        ("GPU：生成计划", 0.85, 2.15, 2.0, BLUE),
        ("CPU：跑测试", 3.00, 2.15, 2.7, YELLOW),
        ("GPU：读结果", 5.85, 2.15, 1.8, BLUE),
        ("CPU：改代码", 7.82, 2.15, 2.4, YELLOW),
        ("GPU：继续推理", 10.38, 2.15, 2.0, BLUE),
    ]
    for label, x, y, w, color in labels:
        box(s, x, y, w, 0.72, fill=color, line=color)
        text(s, label, x + 0.05, y + 0.18, w - 0.1, 0.28, size=12.5, color=BG, bold=True, align=PP_ALIGN.CENTER)
    rule(s, 0.90, 3.18, 12.35, 3.18, color=FAINT, width=1.5)
    text(s, "时间 →", 11.50, 3.25, 0.90, 0.25, size=10, color=FAINT, align=PP_ALIGN.RIGHT)
    card(s, 0.78, 3.72, 3.72, 1.73, "机会", "工具运行时，GPU 可能空出来，可以服务别的任务。", color=GREEN, body_size=14)
    card(s, 4.80, 3.72, 3.72, 1.73, "风险", "工具突然结束，多批请求一起返回，首字延迟可能恶化。", color=CORAL, body_size=14)
    card(s, 8.82, 3.72, 3.72, 1.73, "核心问题", "怎样知道“还有多少工作”，并把这个信息变成安全动作？", color=CYAN, body_size=14)
    lesson(s, "研究对象应是完整任务的交替阶段，而不是孤立的一条 LLM 请求或 shell 命令。")
    add_source(s, "概念口径：analysis/ROADMAP.md；KV = GPU 中保存的对话中间状态")

    # 3 — data
    s = new_slide(prs)
    add_title(s, "我们先补齐了公开数据缺少的东西：原始命令与子命令级资源数据", 3, "数据基础")
    card(s, 0.75, 1.45, 2.90, 2.10, "SQLGlot · 200 tasks", "同一仓库、任务相似\n适合学习重复命令与测试阶段", color=CYAN, body_size=14)
    card(s, 3.83, 1.45, 2.90, 2.10, "PennyLane · 76 tasks", "同一仓库、负载更高\n含有效子命令级内核计数", color=PURPLE, body_size=14)
    card(s, 6.91, 1.45, 2.90, 2.10, "SWE · 377 tasks", "跨 205+ 仓库\n适合验证广度，不适合假装同仓库复用", color=BLUE, body_size=14)
    card(s, 9.99, 1.45, 2.60, 2.10, "TB · 239 tasks", "工具更杂\n适合看分布，不等于高并发实机", color=YELLOW, body_size=14)
    text(s, "每条执行命令保留什么？", 0.82, 3.95, 3.0, 0.38, size=16, bold=True)
    stages = [
        ("原始工具参数", 0.84, 4.50, 2.25, CYAN),
        ("shell 子命令", 3.36, 4.50, 2.15, BLUE),
        ("完整命令 ↔ 子命令", 5.78, 4.50, 2.55, PURPLE),
        ("CPU / 内存 / 读写 / 时长", 8.60, 4.50, 3.15, GREEN),
    ]
    for i, (label, x, y, w, color) in enumerate(stages):
        box(s, x, y, w, 0.72, fill=PANEL2, line=color)
        text(s, label, x + 0.05, y + 0.18, w - 0.1, 0.30, size=12.3, color=color, bold=True, align=PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + w + 0.08, y + 0.27, 0.24, 0.16, fill=FAINT)
    lesson(s, "这批 trace 的价值不是“更多行”，而是能把一句 shell 命令拆成可解释的真实工作。")
    add_source(s, "证据：tool-resource-canonical-objective.md §6–8；这些数据都已用于开发，不是最终独立验证")

    # 4 — CacheWise
    s = new_slide(prs)
    add_title(s, "阶段 1：CacheWise 的整段参数相似度，在我们的 SWE 上反而更差", 4, "从论文复现开始")
    text(s, "CacheWise 的直觉", 0.78, 1.48, 2.4, 0.35, size=16, bold=True)
    text(s, "把完整工具参数转成文本向量 → 聚类 → 用同类历史估计工具还要多久", 0.78, 1.91, 5.25, 0.66, size=14.5, color=MUTED)
    arrow(s, 5.30, 2.03, 0.62, 0.25, fill=CYAN)
    text(s, "我们的结果：越低越好", 6.35, 1.48, 3.2, 0.35, size=16, bold=True)
    bar_chart(
        s,
        ["工具名历史", "C20", "C50", "C100"],
        [0.780, 0.857, 1.179, 1.366],
        6.35,
        1.95,
        5.65,
        2.20,
        colors=[GREEN, BLUE, PURPLE, CORAL],
        suffix=" s",
        max_value=1.50,
    )
    card(s, 0.78, 3.05, 5.10, 1.83, "为什么失败？", "“python -m pytest”说明要跑测试，却不知道仓库规模、依赖是否缓存、测试是否会下载数据。文字像，不等于剩余工作像。", color=CORAL, body_size=13.6)
    stat(s, 6.40, 4.43, 1.75, "+0.587 s", "C100 额外后悔", color=CORAL, note="95% CI +0.120~+1.188")
    stat(s, 8.35, 4.43, 1.75, "46 / 44", "有帮助 / 有伤害", color=YELLOW, note="变化次数几乎对半")
    stat(s, 10.30, 4.43, 1.75, "35.5 s", "最差 1% 后悔", color=CORAL, note="工具名对照 17.3 s")
    lesson(s, "文本只能描述“在做什么”；资源预测还必须描述“做多少、做到哪、环境还差什么”。")
    add_source(s, "证据：analysis/results/cachewise-swe-reproduction-20260731/result.json；论文：arXiv:2606.16824")

    # 5 — same repo predictor
    s = new_slide(prs)
    add_title(s, "阶段 2：同一仓库更可学，收益来自任务结构", 5, "转向真实工作量")
    text(s, "SQLGlot50 · 四目标平均准确率（越高越好）", 0.82, 1.48, 5.4, 0.35, size=16, bold=True)
    bar_chart(
        s,
        ["无脑选最多档", "子命令历史库", "任务感知预测"],
        [70.783, 80.203, 84.463],
        0.80,
        1.95,
        5.65,
        2.15,
        colors=[FAINT, BLUE, GREEN],
        suffix="%",
        max_value=100,
    )
    text(s, "迁移到 PennyLane（越高越好）", 6.90, 1.48, 4.2, 0.35, size=16, bold=True)
    stat(s, 6.90, 1.98, 2.30, "75.680%", "子命令历史库", color=BLUE)
    stat(s, 9.45, 1.98, 2.30, "78.326%", "任务感知预测", color=GREEN, note="+2.646 个百分点")
    text(s, "任务感知预测增加了什么？", 0.86, 4.23, 3.5, 0.35, size=16, bold=True)
    add_bullets(
        s,
        [
            "完整命令历史，而不是只看单个子命令",
            "同一测试套件的第几次运行（环境可能已变热）",
            "pip / pytest 的等价工作表示；没有证据时回退子命令历史库",
        ],
        0.88,
        4.68,
        8.2,
        1.20,
        size=14,
    )
    chip(s, "77 次改对", 9.50, 4.45, 1.25, fill=GREEN, color=BG)
    chip(s, "25 次改错", 10.90, 4.45, 1.35, fill=CORAL, color=BG)
    text(s, "迁移收益来自多个 task，\n但绝对提升仍然有限。", 9.55, 5.05, 2.65, 0.68, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    lesson(s, "同仓库 trace 提供可重复结构；它能改进预测，却没有自动回答“调度该做什么”。")
    add_source(s, "证据：sqlglot50-multitarget-sota-v1；pennylane-multitarget-transfer-validation-v1")

    # 6 — KB architecture
    s = new_slide(prs)
    add_title(s, "阶段 3：换索引结构不是答案，证据才是", 6, "澄清历史库的角色")
    text(s, "评分单位始终是完整命令", 0.80, 1.48, 3.6, 0.34, size=16, bold=True)
    stages = [
        ("完整命令", "pytest -n 8 tests/core", CYAN),
        ("子命令", "拆出可归因的执行片段", BLUE),
        ("历史证据", "精确参数 / 前缀 / 工具名", PURPLE),
        ("四个预测", "时长、CPU、内存、磁盘档位", GREEN),
    ]
    for i, (heading, body, color) in enumerate(stages):
        x = 0.78 + i * 3.05
        box(s, x, 2.05, 2.58, 1.25, fill=PANEL, line=color)
        text(s, heading, x + 0.12, 2.20, 2.34, 0.30, size=14, color=color, bold=True, align=PP_ALIGN.CENTER)
        text(s, body, x + 0.12, 2.62, 2.34, 0.42, size=11.2, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + 2.68, 2.55, 0.26, 0.18, fill=FAINT)
    card(s, 0.82, 3.75, 3.65, 1.56, "子命令的作用", "把复合命令拆细，找到是哪一段产生 CPU、内存、读写和时长。", color=BLUE, body_size=13.1)
    card(s, 4.83, 3.75, 3.65, 1.56, "历史库的作用", "保存已完成任务的因果历史；当前任务和失败任务不能偷看。", color=PURPLE, body_size=13.1)
    card(s, 8.84, 3.75, 3.65, 1.56, "预测器的作用", "把多层证据合成完整命令的档位概率；没有可靠证据时回退。", color=GREEN, body_size=13.1)
    lesson(s, "换数据结构只有在它改变了可用证据、预测和下游动作时才有科研意义。")
    add_source(s, "口径：tool-resource-canonical-objective.md §2–3；当前历史库按完整参数、参数前缀、工具名逐层回退")

    # 7 — semantics
    s = new_slide(prs)
    add_title(s, "阶段 4：因素之间的联系，才能解释真实工作量", 7, "手写规则、离线语言模型与上界")
    card(s, 0.80, 1.45, 3.80, 2.20, "pip：收益上界很小", "即使知道完美的 pip 工作类型，四目标平均准确率最多只提高 0.392 个百分点。\n\n继续雕刻 package 规则不划算。", color=FAINT, body_size=13.6)
    card(s, 4.78, 1.45, 3.80, 2.20, "pytest：两个因素要一起看", "只看并行进程数，高内存档召回率 2.381%。\n\n同时看并行进程数 × 测试范围后，达到 54.762%。", color=GREEN, body_size=13.6)
    card(s, 8.76, 1.45, 3.80, 2.20, "通用 agent：没有捷径", "让 agent 从 trace 自动生成复杂模式，没有得到可区分的新状态；任意代码还带来泄漏与验证成本。", color=CORAL, body_size=13.6)
    text(s, "pytest 的关键不是识别名字，而是估算并行度与工作范围的乘法效应", 0.90, 4.14, 11.5, 0.38, size=16, bold=True, align=PP_ALIGN.CENTER)
    box(s, 2.05, 4.78, 2.2, 0.70, fill=PANEL2, line=BLUE)
    text(s, "并行进程数", 2.05, 4.98, 2.2, 0.25, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text(s, "×", 4.60, 4.88, 0.45, 0.38, size=24, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    box(s, 5.30, 4.78, 2.2, 0.70, fill=PANEL2, line=PURPLE)
    text(s, "测试范围", 5.30, 4.98, 2.2, 0.25, size=14, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    arrow(s, 7.86, 5.00, 0.70, 0.20, fill=FAINT)
    box(s, 8.82, 4.78, 2.45, 0.70, fill=GREEN, line=GREEN)
    text(s, "内存压力", 8.82, 4.98, 2.45, 0.25, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
    lesson(s, "离线 LM 可以抽取少量、可验证的因素；真正的成本仍应由 trace 学习，而不是让 LM 直接猜。")
    add_source(s, "证据：tool-resource-canonical-objective.md §5 Prediction and tool understanding")

    # 8 — continuous update
    s = new_slide(prs)
    add_title(s, "阶段 5：持续观察能提高准确率，但通常发现得太晚", 8, "动态更新不是自动有用")
    text(s, "SQLGlot：时长加权准确率（越高越好）", 0.82, 1.48, 5.3, 0.35, size=16, bold=True)
    bar_chart(
        s,
        ["开始前预测", "完整命令存活时间", "子命令存活时间"],
        [60.260, 75.827, 83.550],
        0.78,
        2.00,
        6.05,
        2.20,
        colors=[FAINT, BLUE, GREEN],
        suffix="%",
        max_value=100,
    )
    text(s, "但它提前多久变得稳定正确？", 7.30, 1.48, 4.5, 0.35, size=16, bold=True)
    stat(s, 7.30, 2.05, 2.15, "49.2 ms", "完整命令提前量", color=BLUE)
    stat(s, 9.72, 2.05, 2.15, "38.2 ms", "子命令提前量", color=GREEN)
    box(s, 7.30, 3.53, 4.57, 0.86, fill=PANEL2, line=CORAL)
    text(s, "调度动作至少需要 500 ms 提前量", 7.50, 3.77, 4.17, 0.30, size=15, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    text(s, "正确得更晚", 2.15, 4.66, 1.60, 0.30, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(s, 3.80, 4.73, 2.40, 0.18, fill=PANEL2)
    text(s, "动作窗口已经过去", 6.35, 4.66, 2.10, 0.30, size=13, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    lesson(s, "预测是否有价值，要同时看“正确多少”和“提前多久”；晚到的正确答案不能指导调度。")
    add_source(s, "证据：sqlglot50-empirical-command-survival-v2/result.json")

    # 9 — action
    s = new_slide(prs)
    add_title(s, "阶段 6：预测准确 ≠ 调度安全；运行时反馈反而更可靠", 9, "从预测转向动作")
    stat(s, 0.80, 1.48, 2.55, "−45.010%", "CPU 预留量", color=GREEN, note="因果运行时反馈；259 tasks")
    stat(s, 3.58, 1.48, 2.55, "+1.543%", "工具自身变慢", color=YELLOW, note="反馈动作的代价")
    stat(s, 6.36, 1.48, 2.55, "−14.620%", "理想 CPU 调度", color=BLUE, note="知道未来的上界")
    stat(s, 9.14, 1.48, 2.55, "−48.823%", "理想内存调度", color=PURPLE, note="时间变化的内存上界")
    text(s, "为什么真实预测器没直接变成安全调度器？", 0.84, 3.05, 6.0, 0.36, size=16, bold=True)
    add_bullets(
        s,
        [
            "峰值 CPU 不是可持续配额：按峰值太保守，按均值又会重叠爆发",
            "内存三档不等于安全上界；少量低估就可能制造容量暴露",
            "某些命令没有可用预测时，硬性资源准入会让全部任务互相等待",
        ],
        0.86,
        3.56,
        7.4,
        1.62,
        size=14,
    )
    card(s, 8.70, 3.08, 3.55, 2.12, "正确的研究顺序", "1. 先证明动作有上界\n2. 再设计不会死锁的反馈控制\n3. 最后问预测器是否增加收益", color=CYAN, body_size=14)
    lesson(s, "调度器需要的是随时间变化的资源轨迹和安全退路，而不是把分类标签当作硬配额。")
    add_source(s, "证据：swe100-277-cpu-feedback-generality-v1；CPU-idle / temporal-RSS oracle results")

    # 10 — tool gap physical
    s = new_slide(prs)
    add_title(s, "阶段 7：工具空档确实能换来约 33% 完成时间收益，但返回尾部变差", 10, "真实 A100 信号")
    text(s, "动作：前台任务进入长工具阶段时，临时放入一个等待任务", 0.82, 1.48, 8.8, 0.38, size=16, bold=True)
    box(s, 0.88, 2.12, 2.20, 0.78, fill=BLUE, line=BLUE)
    text(s, "前台：GPU", 0.88, 2.36, 2.20, 0.26, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
    arrow(s, 3.28, 2.40, 0.68, 0.18, fill=FAINT)
    box(s, 4.16, 2.12, 2.20, 0.78, fill=YELLOW, line=YELLOW)
    text(s, "前台：工具", 4.16, 2.36, 2.20, 0.26, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
    arrow(s, 6.56, 2.40, 0.68, 0.18, fill=GREEN)
    box(s, 7.44, 2.12, 2.25, 0.78, fill=GREEN, line=GREEN)
    text(s, "借入新任务", 7.44, 2.36, 2.25, 0.26, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
    arrow(s, 9.89, 2.40, 0.68, 0.18, fill=CORAL)
    box(s, 10.77, 2.12, 1.75, 0.78, fill=CORAL, line=CORAL)
    text(s, "一起返回", 10.77, 2.36, 1.75, 0.26, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
    stat(s, 0.82, 3.52, 2.55, "−32.748%", "平均完成时间 · r1", color=GREEN)
    stat(s, 3.58, 3.52, 2.55, "−33.102%", "平均完成时间 · r2", color=GREEN)
    stat(s, 6.34, 3.52, 2.55, "1.048×", "最慢 1% 首字 · r1", color=YELLOW, note="预设上限 1.05×")
    stat(s, 9.10, 3.52, 2.55, "1.115×", "最慢 1% 首字 · r2", color=CORAL, note="未通过尾部标准")
    text(s, "静态预测没触发：最高时长档从 30 s 开始，但动作需要 >34.7 s 才值得。", 1.05, 5.12, 11.1, 0.42, size=14.2, color=MUTED, align=PP_ALIGN.CENTER)
    lesson(s, "空档借用是有效动作；下一步不是再预测时长，而是保护前台任务返回时的 GPU 服务。")
    add_source(s, "证据：analysis/results/pennylane-physical-gap-loan-development-v1/result.json")

    # 11 — joint upper bound
    s = new_slide(prs)
    add_title(s, "阶段 8：最大的机会来自 CPU 与 GPU 阶段错开，而不是只调一个资源", 11, "联合调度上界")
    text(s, "PennyLane 70 tasks · 理想上界 · 平均完成时间（越低越好）", 0.82, 1.48, 7.0, 0.36, size=16, bold=True)
    bar_chart(
        s,
        ["只按 GPU 限制", "只按工具限制", "联合看阶段"],
        [79829, 21705, 14100],
        0.82,
        2.00,
        7.10,
        2.25,
        colors=[CORAL, YELLOW, GREEN],
        suffix=" s",
        max_value=85000,
    )
    card(s, 8.40, 1.50, 3.80, 1.35, "不是把并发固定得更大", "联合上界最多同时放入 11 个任务，但每个资源仍不超容量。", color=GREEN, body_size=13.1)
    card(s, 8.40, 3.10, 3.80, 1.35, "而是让不同阶段互补", "GPU 忙时少放推理；GPU 空时让 CPU 工具继续推进。", color=CYAN, body_size=13.1)
    chip(s, "GPU 17.3%", 1.15, 4.85, 1.55, fill=BLUE, color=BG)
    chip(s, "CPU 37.7%", 3.00, 4.85, 1.55, fill=YELLOW, color=BG)
    chip(s, "内存 14.2%", 4.85, 4.85, 1.55, fill=PURPLE, color=BG)
    text(s, "这些仍是 trace 回放上界，不是实机 GPU 结论。", 7.35, 4.88, 4.55, 0.30, size=13.5, color=CORAL, bold=True)
    lesson(s, "真正的 frontier 是跨资源阶段协调：准入、优先级与状态保留必须一起看。")
    add_source(s, "证据：analysis/results/pennylane-joint-phase-packing-v1/result.json；数值为知道未来的理想上界")

    # 12 — closed paths
    s = new_slide(prs)
    add_title(s, "阶段 9：理想上界也没收益，方向才真正关闭", 12, "关闭支线")
    card(s, 0.78, 1.42, 3.75, 1.70, "容器暂停 / 快照", "即使假设暂停与恢复完全免费，平均完成时间仍恶化 2.664%。只释放 3.178% CPU 时间、1.316% 内存时间。", color=CORAL, body_size=12.8)
    card(s, 4.78, 1.42, 3.75, 1.70, "静态资源档位准入", "更准的内存分类仍出现大量容量暴露；两种因果映射方式在 307 s 互相等待。", color=CORAL, body_size=12.8)
    card(s, 8.78, 1.42, 3.75, 1.70, "工具时长优先级", "即使知道真实工具时长，短工具优先也只改善 0.435%；学出来只会更低。", color=FAINT, body_size=12.8)
    card(s, 0.78, 3.52, 3.75, 1.70, "CacheWise / C100", "在当前模拟器中，预测只减少 1.481% 的 KV 重算；知道未来的上界也仅 9.620%。", color=FAINT, body_size=12.8)
    card(s, 4.78, 3.52, 3.75, 1.70, "通用 agent 规则生成", "复杂输出没有创造新的可辨识状态，却增加泄漏、运行成本和验证成本。", color=FAINT, body_size=12.8)
    card(s, 8.78, 3.52, 3.75, 1.70, "本轮 vLLM 原生优先级", "目前没有正式结果：前两组有效，第三组仅辅助 GPU 采样间隔超限。不能从中声称方法失败。", color=YELLOW, body_size=12.8)
    lesson(s, "只有“理想上界也没收益”的方向才真正关闭；框架无效不能当作方法无效。")
    add_source(s, "证据：perfect-container-parking、joint causal admission、revocable lease 与 canonical objective 结果表")

    # 13 — related works
    s = new_slide(prs)
    add_title(s, "三个相关工作已覆盖很多“显而易见”的动作", 13, "Related work：先划清边界")
    card(s, 0.75, 1.42, 3.85, 3.95, "Continuum", "看见工具调用后，为 GPU 中间状态设置保留期限。\n\n期限由工具历史、重算/重载成本和排队代价决定；到期就释放，避免无限占用。\n\n已经覆盖：工具空档中的 KV 保留。", color=CYAN, body_size=13.1)
    card(s, 4.75, 1.42, 3.85, 3.95, "ThunderAgent", "把完整 agent program 作为调度单位，区分推理/工具阶段。\n\n按工具已运行时间降低其 KV 优先级，并做跨 GPU 恢复、迁移和工具环境准备。\n\n已经覆盖：阶段感知 + 状态生命周期。", color=YELLOW, body_size=13.1)
    card(s, 8.75, 1.42, 3.85, 3.95, "Agentix", "截获完整程序的 LLM 调用，按已完成服务量或关键路径调度。\n\n长请求偏向原模型副本保留局部性；短请求发到较空副本。\n\n已经覆盖：程序级优先级 + 到达后的路由。", color=PURPLE, body_size=13.1)
    lesson(s, "“看工具阶段”“保留 KV”“按 program 排队”“KV-local 路由”都不能再作为我们的新贡献。")
    add_source(s, "Continuum arXiv:2511.02230 · ThunderAgent arXiv:2602.13692v3 · Agentix NSDI'26")

    # 14 — overlap
    s = new_slide(prs)
    add_title(s, "我们的可检验差异，只剩“工具内部进度能否让动作更好”", 14, "交叉与候选贡献")
    rows = [
        ("外层工具名 / 已运行时间", "✓", "✓", "—", "已有对照"),
        ("程序级排队 / 到达后路由", "先来先服务", "✓", "✓", "已有对照"),
        ("GPU 中间状态保留 / 迁移", "保留期限", "✓", "局部性", "已有对照"),
        ("shell 子命令已完成到哪里", "—", "—", "—", "候选"),
        ("CPU / 内存 / 读写的实时因果状态", "—", "—", "—", "候选"),
        ("同时决定 GPU 返回与 CPU 工具执行", "—", "部分", "—", "候选"),
    ]
    headers = ["信息 / 动作", "Continuum", "ThunderAgent", "Agentix", "我们的位置"]
    xs = [0.72, 4.60, 6.35, 8.45, 10.15]
    ws = [3.72, 1.60, 1.90, 1.55, 2.35]
    for x, w, header in zip(xs, ws, headers, strict=True):
        box(s, x, 1.45, w, 0.52, fill=PANEL2, line=PANEL2, radius=False)
        text(s, header, x + 0.04, 1.59, w - 0.08, 0.25, size=11.5, bold=True, align=PP_ALIGN.CENTER)
    for i, row in enumerate(rows):
        y = 1.97 + i * 0.65
        fill = PANEL if i % 2 == 0 else BG
        for j, (x, w, value) in enumerate(zip(xs, ws, row, strict=True)):
            box(s, x, y, w, 0.65, fill=fill, line=PANEL2, radius=False)
            color = CYAN if j == 4 and value == "候选" else (MUTED if value in {"—", "已有对照"} else WHITE)
            text(s, value, x + 0.06, y + 0.17, w - 0.12, 0.28, size=11.2, color=color, bold=(j == 0 or color == CYAN), align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
    text(s, "候选假设：比“工具已经跑了多久”更细的内部进度，能更早判断返回窗口与资源压力。", 0.85, 6.03, 11.6, 0.36, size=15, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_source(s, "边界依据：analysis/offline/related-work.md；候选不等于已成立贡献")

    # 15 — smallest test
    s = new_slide(prs)
    add_title(s, "下一步先做一个最小因果实验，不先造新系统", 15, "把候选差异变成可证伪问题")
    text(s, "同一批真实 trace，比较三层信息是否真的改变动作", 0.82, 1.46, 7.2, 0.38, size=16, bold=True)
    arms = [
        ("A", "阶段反馈", "只知道：现在在 LLM 还是工具阶段", FAINT),
        ("B", "ThunderAgent 风格", "再加：工具已运行多久、GPU 状态位置", YELLOW),
        ("C", "工具内部状态", "再加：子命令进度与 CPU/内存/读写反馈", CYAN),
    ]
    for i, (letter, heading, body, color) in enumerate(arms):
        y = 2.05 + i * 1.10
        circle(s, 0.90, y + 0.05, 0.52, fill=color)
        text(s, letter, 0.90, y + 0.15, 0.52, 0.22, size=12, color=BG, bold=True, align=PP_ALIGN.CENTER)
        box(s, 1.65, y, 5.55, 0.70, fill=PANEL, line=color)
        text(s, heading, 1.85, y + 0.17, 1.80, 0.28, size=14, color=color, bold=True)
        text(s, body, 3.60, y + 0.17, 3.35, 0.30, size=12.2, color=MUTED)
        if i < 2:
            arrow(s, 1.06, y + 0.82, 0.18, 0.20, fill=FAINT)
    card(s, 7.72, 1.55, 4.45, 1.40, "先在现有 trace 上做", "问 C 是否比 B 改变了借入、返回保护或 KV 保留决策；不需要 GPU。", color=GREEN, body_size=13.2)
    card(s, 7.72, 3.18, 4.45, 1.40, "只有过门槛才上实机", "至少改变足够多独立任务的动作，并在 trace 回放中同时改善完成时间与返回重叠。", color=YELLOW, body_size=13.2)
    card(s, 7.72, 4.81, 4.45, 1.10, "实机对照不能缺", "Continuum 式保留期限、ThunderAgent 式时间衰减、Agentix 式到达后优先级。", color=PURPLE, body_size=12.5)
    lesson(s, "贡献不是“我们看得更细”，而是这份细粒度状态在相同成本下改变了更好的动作。")
    add_source(s, "下一步需冻结 cohort、动作、成本和 GO/STOP 标准后再读结果；不启动新的昂贵实验")

    # 16 — paper arc
    s = new_slide(prs)
    add_title(s, "如果三个环节都成立，论文故事才完整", 16, "从预测到系统效果")
    steps = [
        ("1", "理解工具", "完整命令 → 子命令 → 真实工作", BLUE),
        ("2", "预测状态", "何时返回、CPU/内存/读写如何变化", CYAN),
        ("3", "连接动作", "借入、返回优先、KV 保留、工具进程放置", YELLOW),
        ("4", "证明收益", "任务更快、尾部不坏、移动成本全计入", GREEN),
    ]
    for i, (number, heading, body, color) in enumerate(steps):
        x = 0.78 + i * 3.05
        circle(s, x + 0.91, 1.55, 0.68, fill=color)
        text(s, number, x + 0.91, 1.75, 0.68, 0.24, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
        box(s, x, 2.52, 2.52, 1.65, fill=PANEL, line=color)
        text(s, heading, x + 0.12, 2.75, 2.28, 0.32, size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
        text(s, body, x + 0.18, 3.24, 2.16, 0.58, size=12.2, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + 2.63, 3.20, 0.28, 0.18, fill=FAINT)
    card(s, 0.82, 4.72, 3.55, 1.15, "可能贡献 A", "普通黑盒 agent trace 也能得到完整命令/子命令级资源状态。", color=BLUE, body_size=12.8)
    card(s, 4.88, 4.72, 3.55, 1.15, "可能贡献 B", "细粒度状态改善跨 CPU/GPU 的实际调度动作。", color=CYAN, body_size=12.8)
    card(s, 8.94, 4.72, 3.55, 1.15, "可能贡献 C", "在真实 agent 负载上同时改善完成时间与尾部。", color=GREEN, body_size=12.8)
    lesson(s, "单独的预测准确率不够；系统贡献必须落在完整任务完成时间、尾部与资源成本上。")
    add_source(s, "当前 A 有成熟基础；B 有 trace 回放上界与物理反馈信号；C 尚未成立")

    # 17 — takeaways
    s = new_slide(prs)
    add_title(s, "带走四句话，然后继续问一个更尖锐的问题", 17, "Takeaways")
    takeaways = [
        ("01", "文本相似，不代表真实工作相似。", CORAL),
        ("02", "同仓库结构能提高预测，但不自动产生调度收益。", BLUE),
        ("03", "反馈和阶段错开已有强信号；硬档位准入容易不安全。", YELLOW),
        ("04", "Continuum、ThunderAgent、Agentix 已覆盖外层阶段与 KV 动作。", PURPLE),
    ]
    for i, (number, value, color) in enumerate(takeaways):
        y = 1.48 + i * 0.92
        text(s, number, 0.85, y + 0.05, 0.62, 0.30, size=15, color=color, bold=True)
        rule(s, 1.52, y + 0.25, 2.05, y + 0.25, color=color, width=2)
        text(s, value, 2.25, y, 9.7, 0.45, size=17, bold=True)
    box(s, 0.82, 5.25, 11.75, 0.90, fill=PANEL2, line=CYAN)
    text(
        s,
        "下一问：工具内部进度是否能比“已运行多久”更早、更准地保护 GPU 返回，同时推进 CPU 工作？",
        1.05,
        5.51,
        11.3,
        0.35,
        size=17,
        color=CYAN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    text(s, "证据入口", 0.85, 6.40, 1.05, 0.22, size=9.8, color=FAINT, bold=True)
    text(
        s,
        "analysis/development/tool-resource-canonical-objective.md  ·  analysis/ROADMAP.md  ·  analysis/offline/related-work.md",
        1.75,
        6.36,
        10.45,
        0.26,
        size=9.2,
        color=MUTED,
    )
    text(
        s,
        "论文：arxiv.org/abs/2511.02230  ·  arxiv.org/abs/2602.13692  ·  usenix.org/conference/nsdi26/presentation/luo",
        1.75,
        6.67,
        10.45,
        0.26,
        size=9.2,
        color=MUTED,
    )

    return prs


if __name__ == "__main__":
    HERE.mkdir(parents=True, exist_ok=True)
    deck = build()
    deck.save(OUT)
    print(f"wrote {OUT} ({len(deck.slides)} slides)")
